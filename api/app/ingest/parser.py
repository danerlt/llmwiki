import re

_DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def parse_to_text(filename: str, content_type: str, data: bytes) -> str:
    """文件字节 → 纯文本。MD/TXT 直接解码；PDF→PDFium；DOCX→python-docx；HTML→去标签。"""
    lower = filename.lower()
    if lower.endswith(".pdf") or content_type == "application/pdf":
        return _parse_pdf(data)
    if lower.endswith(".docx") or content_type == _DOCX_CT:
        return _parse_docx(data)
    if lower.endswith((".html", ".htm")) or content_type == "text/html":
        return _parse_html(data.decode("utf-8", errors="replace"))
    return data.decode("utf-8", errors="replace")


def _parse_docx(data: bytes) -> str:
    """DOCX 文本提取：段落 + 表格单元格（python-docx）。"""
    import io

    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _parse_html(html: str) -> str:
    """HTML → 纯文本：标准库去标签，跳过 script/style，零额外依赖。"""
    from html.parser import HTMLParser

    class _Extract(HTMLParser):
        def __init__(self) -> None:
            super().__init__()
            self.parts: list[str] = []
            self._skip = False

        def handle_starttag(self, tag: str, attrs: list) -> None:
            if tag in ("script", "style"):
                self._skip = True

        def handle_endtag(self, tag: str) -> None:
            if tag in ("script", "style"):
                self._skip = False

        def handle_data(self, data: str) -> None:
            if not self._skip and data.strip():
                self.parts.append(data.strip())

    ex = _Extract()
    ex.feed(html)
    return "\n".join(ex.parts)


def _parse_pdf(data: bytes) -> str:
    """PDF 文本层提取：优先 liteparse(PDFium，比 pypdf 更鲁棒、容错无 xref)，
    异常时回退 pypdf。MVP 仅取文本层，OCR 关闭（扫描件 OCR 属 V2）。"""
    try:
        from liteparse import LiteParse

        return LiteParse(ocr_enabled=False).parse(data).text
    except Exception:  # noqa: BLE001 — liteparse 失败则降级到 pypdf，仍失败由上层置 failed
        import io

        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)


def slugify(text: str) -> str:
    """生成 slug：去首尾空白，ASCII 转小写，非字母数字（CJK 等 \\w 字符保留）转连字符。
    若归一化后为空（纯标点/符号标题），按原文 hash 兜底，保证非空且不同原文不撞同一 slug。"""
    raw = text
    text = text.strip().lower()
    text = re.sub(r"[^\w]+", "-", text).strip("-")
    if not text:
        import hashlib

        text = "page-" + hashlib.sha1(raw.strip().encode("utf-8")).hexdigest()[:8]
    return text
